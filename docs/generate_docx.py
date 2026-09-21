import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_report():
    doc = Document()

    # Set margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Style defaults
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Cordia New'
    font.size = Pt(14)
    font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title.add_run("รายงานโครงสร้างและการทำงานของระบบ Backend\n(LUMA AI System Architecture & Data Flow Report)")
    run_title.font.name = 'Cordia New'
    run_title.font.size = Pt(22)
    run_title.bold = True
    run_title.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    doc.add_paragraph()

    # Section 1
    h1 = doc.add_heading(level=1)
    r1 = h1.add_run("1. ภาพรวมสถาปัตยกรรมระบบ (System Architecture Overview)")
    r1.font.name = 'Cordia New'
    r1.font.size = Pt(18)
    r1.bold = True
    r1.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    p = doc.add_paragraph("ระบบ Backend ของโครงการ LUMA พัฒนาขึ้นด้วยภาษา Python โดยใช้ Flask Framework ซึ่งได้รับการออกแบบตามหลักสถาปัตยกรรม Application Factory Pattern และแยกกลุ่มเส้นทางการทำงานออกเป็น Flask Blueprints เพื่อให้ง่ายต่อการดูแลรักษาและขยายระบบในอนาคต")
    p.paragraph_format.line_spacing = 1.15

    # Architecture bullet points
    p = doc.add_paragraph(style='List Bullet')
    r = p.add_run("Framework: ")
    r.bold = True
    p.add_run("Python 3.11 + Flask 3.0.3 (เชื่อมต่อผ่าน WSGI และ Nginx Reverse Proxy)")

    p = doc.add_paragraph(style='List Bullet')
    r = p.add_run("Database Layer: ")
    r.bold = True
    p.add_run("SQLite (database/database.db) เชื่อมต่อและบริหารจัดการผ่าน Flask-SQLAlchemy (ORM)")

    p = doc.add_paragraph(style='List Bullet')
    r = p.add_run("Authentication & Security: ")
    r.bold = True
    p.add_run("JSON Web Token (JWT HS256, อายุ 24 ชม.), Password Hashing (PBKDF2/SHA256), และระบบ Anti-Hopping / IDOR Protection ป้องกันการสวมรอยข้ามบัญชี")

    p = doc.add_paragraph(style='List Bullet')
    r = p.add_run("AI Integration: ")
    r.bold = True
    p.add_run("เชื่อมต่อไปยัง Stable Diffusion AI Server ผ่าน REST API พร้อมระบบ Concurrency Lock (Mutex Queue) ป้องกันงานชนกันบน GPU")

    doc.add_paragraph()

    # Section 2
    h2 = doc.add_heading(level=1)
    r2 = h2.add_run("2. หน้าที่และความรับผิดชอบของแต่ละไฟล์ใน Backend")
    r2.font.name = 'Cordia New'
    r2.font.size = Pt(18)
    r2.bold = True
    r2.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    table_data = [
        ("โฟลเดอร์ / ชื่อไฟล์", "หน้าที่และความรับผิดชอบหลัก"),
        ("backend/run.py", "จุดเริ่มต้นรันเซิร์ฟเวอร์ (Entry Point) ดึงค่าพอร์ต (PORT=5000) และเปิดบริการบน host 0.0.0.0"),
        ("backend/config.py", "ศูนย์กลางคอนฟิกูเรชัน (Central Config) เช่น Secret Key, CORS Origins, URL ของ AI Server"),
        ("backend/app/__init__.py", "Application Factory เริ่มต้น Flask, ผูก SQLAlchemy, ลงทะเบียน Blueprints ทั้งหมด และ Auto-Migration ตารางใน DB"),
        ("backend/app/extensions.py", "ประกาศตัวแปรกลาง db = SQLAlchemy() เพื่อป้องกันปัญหา Circular Import"),
        ("backend/app/models/user.py", "โมเดลตาราง users (id, email, password_hash, created_at) และผูก Cascade Delete ไปยังตารางภาพ"),
        ("backend/app/models/generation.py", "โมเดลตาราง generations รองรับประวัติ 2 หมวดหมู่ (sd_generate และ image_filter) และเก็บไฟล์ภาพเป็น BLOB"),
        ("backend/app/middleware/jwt_auth.py", "มิดเดิลแวร์ @token_required ตรวจสอบ Bearer JWT Token, ถอดรหัส user_id และป้องกันการข้ามสิทธิ์"),
        ("backend/app/routes/auth.py", "เส้นทาง /register (สมัครสมาชิก, ตรวจเมลซ้ำ, แฮชรหัสผ่าน) และ /login (ตรวจสอบรหัสผ่าน, ออก JWT Token)"),
        ("backend/app/routes/sd.py", "เส้นทาง /generate (สร้างภาพ AI), /estimate (ประเมินเวลา), /images/:id (สตรีมภาพ PNG พร้อมเช็กเจ้าของ)"),
        ("backend/app/routes/history.py", "เส้นทาง /history (ดึงประวัติพร้อม Pagination & Filter), ดูรายละเอียดภาพ, และลบประวัติของตนเอง"),
        ("backend/app/routes/profile.py", "เส้นทาง /profile ดึงข้อมูลส่วนตัวผู้ใช้ และสรุปยอดจำนวนภาพที่สร้างแยกตามประเภท"),
        ("backend/app/services/ai_client.py", "ตัวเชื่อมต่อ AI Server มีระบบ Concurrency Queue (GPU Lock) และระบบคำนวณเวลาโดยประมาณ"),
        ("backend/app/utils/error_codes.py", "ฟังก์ชันห่อ Response Format (success_response, error_response) และนิยาม Error Codes มาตรฐาน")
    ]

    table = doc.add_table(rows=len(table_data), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    for row_idx, row in enumerate(table_data):
        for col_idx, text in enumerate(row):
            cell = table.cell(row_idx, col_idx)
            cell.text = text
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            run = p.runs[0]
            run.font.name = 'Cordia New'
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)

            if row_idx == 0:
                run.bold = True
                run.font.size = Pt(14)
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                set_cell_background(cell, "1F4E79")
            else:
                run.font.size = Pt(13)
                if col_idx == 0:
                    run.bold = True
                    set_cell_background(cell, "F2F2F2")
                else:
                    if row_idx % 2 == 0:
                        set_cell_background(cell, "FAFAFA")

    # Set column widths
    for row in table.rows:
        row.cells[0].width = Inches(2.2)
        row.cells[1].width = Inches(4.3)

    doc.add_paragraph()

    # Section 3
    h3 = doc.add_heading(level=1)
    r3 = h3.add_run("3. เจาะลึกกระบวนการทำงานและเส้นทางข้อมูล (Step-by-Step Data Flow)")
    r3.font.name = 'Cordia New'
    r3.font.size = Pt(18)
    r3.bold = True
    r3.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    # 3.1 Register Flow
    h3_1 = doc.add_heading(level=2)
    r3_1 = h3_1.add_run("3.1 การสมัครสมาชิก (Register Flow)")
    r3_1.font.name = 'Cordia New'
    r3_1.font.size = Pt(16)
    r3_1.bold = True
    r3_1.font.color.rgb = RGBColor(0x2F, 0x55, 0x97)

    steps_reg = [
        ("1. ผู้ใช้กรอกข้อมูลบน Frontend: ", "กรอก email และ password แล้วกดปุ่ม Register จากนั้น Frontend ส่งคำขอ POST /api/v1/register พร้อม Body JSON {'email': 'peem@example.com', 'password': 'Password123'}"),
        ("2. Backend Validation: ", "ฟังก์ชัน auth.py รับข้อมูล ตรวจสอบรูปแบบอีเมลด้วย Regex และตรวจความยาวรหัสผ่าน (ต้อง >= 8 ตัวอักษร)"),
        ("3. Duplicate Check (ตรวจอีเมลซ้ำ): ", "ค้นหาในตาราง users ด้วยคำสั่ง User.query.filter_by(email=email).first() หากพบว่ามีอยู่แล้ว จะตอบกลับ HTTP 409 EMAIL_EXISTS"),
        ("4. Password Hashing: ", "นำรหัสผ่านไปผ่านฟังก์ชัน generate_password_hash() เพื่อแปลงเป็นแฮช PBKDF2/SHA256 ที่ปลอดภัย ไม่เก็บรหัสผ่านจริง"),
        ("5. บันทึกลงฐานข้อมูล: ", "สร้าง User object ใหม่แล้วบันทึก (db.session.add(new_user) -> db.session.commit())"),
        ("6. ส่ง Response กลับ Frontend: ", "ส่งรหัสสถานะ HTTP 201 Created พร้อมข้อมูลผู้ใช้ในรูปแบบ JSON เพื่อให้ Frontend พาไปยังหน้า Login")
    ]
    for bold_text, normal_text in steps_reg:
        p = doc.add_paragraph(style='List Bullet')
        r = p.add_run(bold_text)
        r.bold = True
        p.add_run(normal_text)

    doc.add_paragraph()

    # 3.2 Login Flow
    h3_2 = doc.add_heading(level=2)
    r3_2 = h3_2.add_run("3.2 การเข้าสู่ระบบและออก JWT Token (Login Flow)")
    r3_2.font.name = 'Cordia New'
    r3_2.font.size = Pt(16)
    r3_2.bold = True
    r3_2.font.color.rgb = RGBColor(0x2F, 0x55, 0x97)

    steps_login = [
        ("1. ผู้ใช้ส่งคำขอ Login: ", "Frontend ส่งคำขอ POST /api/v1/login พร้อม email และ password"),
        ("2. ตรวจสอบบัญชีและรหัสผ่าน: ", "ค้นหา User ด้วย email หากพบจะนำรหัสผ่านไปตรวจกับค่า Hash ด้วย check_password_hash(user.password_hash, password) หากไม่ตรงจะตอบกลับ HTTP 401 INVALID_CREDENTIALS"),
        ("3. สร้าง JWT Access Token: ", "สร้าง Payload บรรจุ {'user_id': user.id, 'email': user.email, 'exp': now + 24 ชั่วโมง} และเข้ารหัสด้วย Secret Key (HS256)"),
        ("4. ส่ง Token กลับ Frontend: ", "ส่ง JSON ตอบกลับพร้อม access_token ให้ Frontend จัดเก็บใน localStorage สำหรับใช้แนบกับคำขออื่นๆ")
    ]
    for bold_text, normal_text in steps_login:
        p = doc.add_paragraph(style='List Bullet')
        r = p.add_run(bold_text)
        r.bold = True
        p.add_run(normal_text)

    doc.add_paragraph()

    # 3.3 Generation Flow
    h3_3 = doc.add_heading(level=2)
    r3_3 = h3_3.add_run("3.3 การสร้างรูปภาพด้วย AI (Image Generation Flow)")
    r3_3.font.name = 'Cordia New'
    r3_3.font.size = Pt(16)
    r3_3.bold = True
    r3_3.font.color.rgb = RGBColor(0x2F, 0x55, 0x97)

    steps_gen = [
        ("1. Frontend สั่งสร้างภาพ: ", "ส่งคำขอ POST /api/v1/generate พร้อม Header 'Authorization: Bearer <Token>' และพารามิเตอร์ Prompt, Size, Steps, Sampler"),
        ("2. Middleware ตรวจสอบสิทธิ์: ", "มิดเดิลแวร์ @token_required ถอดรหัส Token และสกัด user_id ของผู้ใช้ปัจจุบันเก็บไว้ใน request.user_id"),
        ("3. Concurrency Queue (GPU Lock): ", "คำขอเข้าสู่คิว ai_client.py ซึ่งมี Mutex Lock (_gpu_lock.acquire()) เพื่อให้ GPU ทำงานทีละ 1 งานอย่างปลอดภัย ไม่เกิดปัญหาคิวชนกัน"),
        ("4. ส่งคำขอไปยัง AI Server: ", "ยิง HTTP POST ไปที่ Stable Diffusion WebUI API (/sdapi/v1/txt2img) รอรับผลลัพธ์ภาพแบบ Base64 String"),
        ("5. ปลดล็อกคิว GPU: ", "สั่ง _gpu_lock.release() ทันทีเพื่อให้คำขอถัดไปในคิวเริ่มทำงานต่อได้"),
        ("6. บันทึกรูปภาพลง SQLite (BLOB): ", "แปลง Base64 เป็น Binary Bytes แล้วบันทึกลงตาราง generations ผูกกับ user_id และระบุ category='sd_generate'"),
        ("7. ตอบกลับผลลัพธ์: ", "ส่ง JSON ตอบกลับพร้อม generation_id และ URL /api/v1/images/:id ให้ Frontend นำไปแสดงผล")
    ]
    for bold_text, normal_text in steps_gen:
        p = doc.add_paragraph(style='List Bullet')
        r = p.add_run(bold_text)
        r.bold = True
        p.add_run(normal_text)

    doc.add_paragraph()

    # 3.4 IDOR Protection
    h3_4 = doc.add_heading(level=2)
    r3_4 = h3_4.add_run("3.4 การเปิดดูรูปภาพและระบบป้องกันการข้ามสิทธิ์ (IDOR / Anti-Hopping Protection)")
    r3_4.font.name = 'Cordia New'
    r3_4.font.size = Pt(16)
    r3_4.bold = True
    r3_4.font.color.rgb = RGBColor(0x2F, 0x55, 0x97)

    steps_idor = [
        ("1. การร้องขอรูปภาพ: ", "Frontend สั่ง GET /api/v1/images/:id พร้อมแนบ Token ของผู้ใช้"),
        ("2. การตรวจสอบความเป็นเจ้าของ: ", "Backend ค้นหาภาพในตาราง generations ด้วยเงื่อนไขคู่: id = :id AND user_id = request.user_id เสมอ"),
        ("3. กรณีเป็นเจ้าของภาพ: ", "Backend สตรีมข้อมูล Binary BLOB ในรูปแบบ Content-Type: image/png กลับไปให้ Frontend แสดงผลได้ทันที"),
        ("4. กรณีพยายามแอบดูภาพของผู้อื่น: ", "คำสั่งค้นหาจะไม่พบข้อมูล (เนื่องจาก user_id ไม่ตรงกัน) และระบบจะตอบกลับ HTTP 404 GENERATION_NOT_FOUND ทันที โดยไม่เปิดเผยข้อมูลให้ผู้ใช้คนอื่นทราบ")
    ]
    for bold_text, normal_text in steps_idor:
        p = doc.add_paragraph(style='List Bullet')
        r = p.add_run(bold_text)
        r.bold = True
        p.add_run(normal_text)

    doc.add_paragraph()

    # Section 4
    h4 = doc.add_heading(level=1)
    r4 = h4.add_run("4. สรุปรายการ API Endpoints ทั้งหมดในระบบ")
    r4.font.name = 'Cordia New'
    r4.font.size = Pt(18)
    r4.bold = True
    r4.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    api_table_data = [
        ("Method", "Endpoint Path", "Token?", "คำอธิบายหน้าที่การทำงาน"),
        ("POST", "/api/v1/register", "ไม่จำเป็น", "สมัครสมาชิกผู้ใช้งานใหม่ (รับ email, password)"),
        ("POST", "/api/v1/login", "ไม่จำเป็น", "เข้าสู่ระบบ ตรวจสอบรหัสผ่าน และรับ JWT Token"),
        ("GET", "/api/v1/models", "ไม่จำเป็น", "ดึงรายชื่อ Checkpoint Models ทั้งหมดจาก AI Server"),
        ("POST", "/api/v1/estimate", "ไม่จำเป็น", "คำนวณเวลาประเมินล่วงหน้าในการสร้างภาพ"),
        ("POST", "/api/v1/generate", "ต้องมี", "สั่งสร้างรูปภาพ AI (ผูกข้อมูลกับเจ้าของบัญชี)"),
        ("GET", "/api/v1/images/:id", "ต้องมี", "สตรีมไฟล์รูปภาพ Binary PNG (จำกัดสิทธิ์เฉพาะเจ้าของ)"),
        ("GET", "/api/v1/history", "ต้องมี", "ดึงรายการประวัติภาพพร้อม Pagination และ Filter"),
        ("GET", "/api/v1/history/:id", "ต้องมี", "ดึงรายละเอียดพารามิเตอร์เต็มของภาพรายการนั้น"),
        ("DELETE", "/api/v1/history/:id", "ต้องมี", "ลบประวัติรูปภาพของตนเองออกจากระบบ"),
        ("GET", "/api/v1/profile", "ต้องมี", "ดึงข้อมูลโปรไฟล์และสรุปสถิติจำนวนภาพที่สร้างทั้งหมด")
    ]

    t_api = doc.add_table(rows=len(api_table_data), cols=4)
    t_api.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_api.autofit = False

    for row_idx, row in enumerate(api_table_data):
        for col_idx, text in enumerate(row):
            cell = t_api.cell(row_idx, col_idx)
            cell.text = text
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            run = p.runs[0]
            run.font.name = 'Cordia New'
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)

            if row_idx == 0:
                run.bold = True
                run.font.size = Pt(14)
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                set_cell_background(cell, "1F4E79")
            else:
                run.font.size = Pt(13)
                if col_idx == 0:
                    run.bold = True
                if row_idx % 2 == 0:
                    set_cell_background(cell, "FAFAFA")

    t_api.rows[0].cells[0].width = Inches(1.0)
    t_api.rows[0].cells[1].width = Inches(2.0)
    t_api.rows[0].cells[2].width = Inches(1.0)
    t_api.rows[0].cells[3].width = Inches(2.5)

    doc.add_paragraph()

    # Section 5
    h5 = doc.add_heading(level=1)
    r5 = h5.add_run("5. มาตรฐานรหัสข้อผิดพลาด (Standard Error Codes)")
    r5.font.name = 'Cordia New'
    r5.font.size = Pt(18)
    r5.bold = True
    r5.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    err_table_data = [
        ("Error Code", "HTTP Status", "คำอธิบายสาเหตุ"),
        ("VALIDATION_ERROR", "400", "ข้อมูลที่ส่งมาไม่ถูกต้องตามเงื่อนไข (เช่น อีเมลผิดฟอร์แมต, รหัสผ่านสั้นกว่า 8 ตัว)"),
        ("INVALID_CREDENTIALS", "401", "อีเมลหรือรหัสผ่านไม่ถูกต้อง"),
        ("UNAUTHORIZED", "401", "ไม่มี Token หรือ Token หมดอายุ / ลายเซ็นไม่ถูกต้อง"),
        ("EMAIL_EXISTS", "409", "อีเมลนี้ถูกลงทะเบียนในระบบแล้ว (ตรวจพบชื่อซ้ำ)"),
        ("GENERATION_NOT_FOUND", "404", "ไม่พบเรคคอร์ดรูปภาพ หรือพยายามเข้าถึงภาพของผู้อื่น"),
        ("IMAGE_NOT_FOUND", "404", "ไม่พบข้อมูลรูปภาพแบบ Binary ในฐานข้อมูล"),
        ("AI_SERVER_BUSY", "409", "คิวงานบน AI Server แน่นเกินเวลาที่กำหนดให้รอได้"),
        ("AI_SERVER_TIMEOUT", "504", "AI Server ประมวลผลช้าเกินเวลาที่กำหนด (เกิน 75 วินาที)"),
        ("AI_SERVER_ERROR", "503", "เกิดข้อผิดพลาดฝั่ง AI Server หรือไม่สามารถเชื่อมต่อได้"),
        ("INTERNAL_SERVER_ERROR", "500", "เกิดข้อผิดพลาดที่ไม่คาดคิดภายในระบบเซิร์ฟเวอร์ Backend")
    ]

    t_err = doc.add_table(rows=len(err_table_data), cols=3)
    t_err.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_err.autofit = False

    for row_idx, row in enumerate(err_table_data):
        for col_idx, text in enumerate(row):
            cell = t_err.cell(row_idx, col_idx)
            cell.text = text
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            run = p.runs[0]
            run.font.name = 'Cordia New'
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)

            if row_idx == 0:
                run.bold = True
                run.font.size = Pt(14)
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                set_cell_background(cell, "1F4E79")
            else:
                run.font.size = Pt(13)
                if col_idx == 0:
                    run.bold = True
                if row_idx % 2 == 0:
                    set_cell_background(cell, "FAFAFA")

    # Output path
    output_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "docs")
    os.makedirs(output_dir, exist_ok=True)
    output_path = os.path.join(output_dir, "BACKEND_REPORT.docx")
    doc.save(output_path)
    print(f"Successfully generated DOCX at: {output_path}")

if __name__ == "__main__":
    create_report()
